import json
import os
import tempfile
import unittest
from datetime import date, datetime

from calendardesk import model
from calendardesk.store import Store


def task(**kw):
    t = model.new_task(kw.pop("title", "t"), kw.pop("date", date(2026, 1, 1)))
    t.update(kw)
    return t


class RecurrenceTests(unittest.TestCase):
    def test_one_off_only_on_its_date(self):
        t = task(date="2026-03-10")
        self.assertTrue(model.occurs_on(t, date(2026, 3, 10)))
        self.assertFalse(model.occurs_on(t, date(2026, 3, 11)))

    def test_daily_interval(self):
        t = task(date="2026-03-10", repeat="daily", interval=2)
        self.assertTrue(model.occurs_on(t, date(2026, 3, 12)))
        self.assertFalse(model.occurs_on(t, date(2026, 3, 11)))
        self.assertFalse(model.occurs_on(t, date(2026, 3, 8)))

    def test_weekly_weekdays_and_interval(self):
        # 2026-03-09 is a Monday
        t = task(date="2026-03-09", repeat="weekly", weekdays=[0, 2], interval=2)
        self.assertTrue(model.occurs_on(t, date(2026, 3, 11)))   # Wed, week 0
        self.assertFalse(model.occurs_on(t, date(2026, 3, 16)))  # Mon, week 1
        self.assertTrue(model.occurs_on(t, date(2026, 3, 23)))   # Mon, week 2

    def test_weekly_default_weekday(self):
        t = task(date="2026-03-11", repeat="weekly")
        self.assertTrue(model.occurs_on(t, date(2026, 3, 18)))
        self.assertFalse(model.occurs_on(t, date(2026, 3, 19)))

    def test_monthly_clamps_to_month_end(self):
        t = task(date="2026-01-31", repeat="monthly")
        self.assertTrue(model.occurs_on(t, date(2026, 2, 28)))
        self.assertTrue(model.occurs_on(t, date(2026, 4, 30)))
        self.assertFalse(model.occurs_on(t, date(2026, 4, 29)))

    def test_yearly_feb29(self):
        t = task(date="2024-02-29", repeat="yearly")
        self.assertTrue(model.occurs_on(t, date(2025, 2, 28)))
        self.assertTrue(model.occurs_on(t, date(2028, 2, 29)))
        self.assertFalse(model.occurs_on(t, date(2028, 2, 28)))

    def test_until_and_skip(self):
        t = task(date="2026-03-01", repeat="daily", until="2026-03-05",
                 skip_dates=["2026-03-03"])
        self.assertTrue(model.occurs_on(t, date(2026, 3, 5)))
        self.assertFalse(model.occurs_on(t, date(2026, 3, 6)))
        self.assertFalse(model.occurs_on(t, date(2026, 3, 3)))

    def test_done_per_occurrence(self):
        t = task(date="2026-03-01", repeat="daily")
        model.set_done(t, date(2026, 3, 2), True)
        self.assertTrue(model.is_done(t, date(2026, 3, 2)))
        self.assertFalse(model.is_done(t, date(2026, 3, 3)))
        model.set_done(t, date(2026, 3, 2), False)
        self.assertFalse(model.is_done(t, date(2026, 3, 2)))


class DayListTests(unittest.TestCase):
    def test_completed_sort_last_then_time(self):
        a = task(title="a", date="2026-03-01", time="10:00")
        b = task(title="b", date="2026-03-01", time="08:00", done=True)
        c = task(title="c", date="2026-03-01", time="09:00")
        d = task(title="d", date="2026-03-01")
        got = [t["title"] for t in model.tasks_on([a, b, c, d], date(2026, 3, 1))]
        self.assertEqual(got, ["c", "a", "d", "b"])

    def test_rollover_moves_unfinished_one_offs_only(self):
        old = task(title="old", date="2026-03-01")
        done = task(title="done", date="2026-03-01", done=True)
        rec = task(title="rec", date="2026-03-01", repeat="daily")
        moved = model.rollover([old, done, rec], date(2026, 3, 4))
        self.assertEqual(moved, 1)
        self.assertEqual(old["date"], "2026-03-04")
        self.assertEqual(old["moved_from"], "2026-03-01")
        self.assertEqual(done["date"], "2026-03-01")
        self.assertEqual(rec["date"], "2026-03-01")


class CountdownTests(unittest.TestCase):
    def test_one_off(self):
        c = {"date": "2026-03-10", "yearly": False}
        self.assertEqual(model.countdown_info(c, date(2026, 3, 1)), (9, date(2026, 3, 10), None))
        self.assertEqual(model.countdown_info(c, date(2026, 3, 12))[0], -2)

    def test_yearly_next_occurrence_and_years(self):
        c = {"date": "2020-03-05", "yearly": True}
        days, nxt, years = model.countdown_info(c, date(2026, 3, 6))
        self.assertEqual(nxt, date(2027, 3, 5))
        self.assertEqual(years, 7)
        self.assertEqual(model.countdown_info(c, date(2026, 3, 5)), (0, date(2026, 3, 5), 6))


class SchoolWeekTests(unittest.TestCase):
    def test_week_number(self):
        terms = [{"name": "T1", "start": "2026-06-10", "end": "2026-07-31"}]
        self.assertIsNone(model.school_week(terms, date(2026, 6, 7)))
        self.assertEqual(model.school_week(terms, date(2026, 6, 10)), ("T1", 1))
        self.assertEqual(model.school_week(terms, date(2026, 6, 15)), ("T1", 2))
        self.assertIsNone(model.school_week(terms, date(2026, 8, 3)))


class ThemeTests(unittest.TestCase):
    def test_mix_and_readable(self):
        from calendardesk import theme
        self.assertEqual(theme.mix("#ffffff", "#000000", 1), "#ffffff")
        self.assertEqual(theme.mix("#ffffff", "#000000", 0), "#000000")
        self.assertEqual(theme.readable_on("#ffffff"), "#111111")
        self.assertEqual(theme.readable_on("#000000"), "#ffffff")

    def test_palette_overrides(self):
        from calendardesk import theme
        p = theme.palette({"theme": "paper", "accent": "#123456", "colors": {"weekend": "#00ff00"}})
        self.assertEqual(p["accent"], "#123456")
        self.assertEqual(p["weekend"], "#00ff00")
        self.assertEqual(p["bg"], theme.PRESETS["paper"]["bg"])
        self.assertEqual(theme.palette({"theme": "nope"})["bg"], theme.PRESETS["dark"]["bg"])


class StoreTests(unittest.TestCase):
    def test_roundtrip_and_backup(self):
        with tempfile.TemporaryDirectory() as d:
            s = Store(d)
            s.data["tasks"].append(task(title="x"))
            s.save()
            s.backup(today=date(2026, 3, 1))
            s2 = Store(d)
            self.assertEqual(s2.data["tasks"][0]["title"], "x")
            self.assertIn("settings", s2.data)
            self.assertTrue(os.path.exists(os.path.join(d, "backups", "data-2026-03-01.json")))

    def test_corrupt_file_is_preserved(self):
        with tempfile.TemporaryDirectory() as d:
            with open(os.path.join(d, "data.json"), "w") as f:
                f.write("{not json")
            s = Store(d)
            self.assertEqual(s.data["tasks"], [])
            self.assertTrue(any(n.startswith("data.corrupt-") for n in os.listdir(d)))

    def test_save_syncs_before_replace(self):
        """EXTRA-2: the temp file reaches the disk before it replaces data.json."""
        from unittest import mock
        calls = []
        real_replace = os.replace
        with tempfile.TemporaryDirectory() as d:
            s = Store(d)
            s.data["tasks"].append(task(title="x"))
            with mock.patch("calendardesk.store.os.fsync", side_effect=lambda fd: calls.append("fsync")),                     mock.patch("calendardesk.store.os.replace",
                               side_effect=lambda a, b: (calls.append("replace"), real_replace(a, b))):
                s.save()
            self.assertEqual(calls, ["fsync", "replace"])
            self.assertEqual(Store(d).data["tasks"][0]["title"], "x")

    def test_prune_reminder_state(self):
        with tempfile.TemporaryDirectory() as d:
            s = Store(d)
            s.data["reminders"] = {"a|2026-03-01": "fired", "b|2026-03-09": "fired"}
            s.prune_reminders(date(2026, 3, 10))
            self.assertEqual(list(s.data["reminders"]), ["b|2026-03-09"])


if __name__ == "__main__":
    unittest.main()


class ChecklistTests(unittest.TestCase):
    def _task(self, n=3, **kw):
        t = model.new_task("Report", date(2026, 3, 2))
        t.update(kw)
        t["items"] = [model.new_item(f"step {i}") for i in range(n)]
        return t

    def test_one_off_items(self):
        d = date(2026, 3, 2)
        t = self._task()
        a, b, c = t["items"]
        model.set_item_done(t, a, d, True)
        model.set_item_done(t, b, d, True)
        self.assertEqual(model.progress(t, d), (2, 3))
        self.assertFalse(model.is_done(t, d))
        model.set_item_done(t, c, d, True)
        self.assertTrue(model.is_done(t, d))
        model.set_item_done(t, a, d, False)
        self.assertFalse(model.is_done(t, d))
        self.assertEqual(model.progress(t, d), (2, 3))

    def test_task_tick_ticks_items_untick_keeps(self):
        d = date(2026, 3, 2)
        t = self._task()
        model.set_task_done(t, d, True)
        self.assertTrue(all(model.item_done(t, i, d) for i in t["items"]))
        model.set_task_done(t, d, False)
        self.assertTrue(all(model.item_done(t, i, d) for i in t["items"]))
        self.assertFalse(model.is_done(t, d))

    def test_repeating_items_per_day(self):
        d1, d2 = date(2026, 3, 2), date(2026, 3, 3)
        t = self._task(n=2, repeat="daily")
        model.set_item_done(t, t["items"][0], d1, True)
        self.assertTrue(model.item_done(t, t["items"][0], d1))
        self.assertFalse(model.item_done(t, t["items"][0], d2))
        self.assertFalse(model.is_done(t, d1))
        model.set_item_done(t, t["items"][1], d1, True)
        self.assertTrue(model.is_done(t, d1))
        self.assertFalse(model.is_done(t, d2))
        self.assertEqual(model.progress(t, d2), (0, 2))

    def test_old_task_without_items(self):
        d = date(2026, 3, 2)
        t = model.new_task("Old", d)
        t.pop("items", None)
        self.assertEqual(model.progress(t, d), (0, 0))
        model.set_task_done(t, d, True)
        self.assertTrue(model.is_done(t, d))


class ExportTests(unittest.TestCase):
    def test_checklist_column(self):
        from openpyxl import load_workbook
        from calendardesk.export import export_xlsx
        d = date(2026, 3, 2)
        t = model.new_task("Report", d)
        t["items"] = [model.new_item(x) for x in "abc"]
        model.set_item_done(t, t["items"][0], d, True)
        data = {"tasks": [t], "countdowns": [], "notes": []}
        path = os.path.join(tempfile.mkdtemp(), "out.xlsx")
        export_xlsx(data, path, d)
        ws = load_workbook(path)["Tasks"]
        header = [c.value for c in ws[1]]
        col = header.index("Checklist")
        self.assertEqual(ws[2][col].value, "1/3: a ✓; b; c")

    def test_text_starting_with_equals_stays_text(self):
        """BUG-8: titles and notes are never written as formulas."""
        from openpyxl import load_workbook
        from calendardesk.export import export_xlsx
        d = date(2026, 3, 2)
        t = model.new_task('=HYPERLINK("http://x","click")', d)
        t["note"] = "=cmd|' /C calc'!A0"
        data = {"tasks": [t], "countdowns": [], "notes": [{"text": "=2+3"}]}
        path = os.path.join(tempfile.mkdtemp(), "out.xlsx")
        export_xlsx(data, path, d)
        wb = load_workbook(path)
        ws = wb["Tasks"]
        header = [c.value for c in ws[1]]
        for name, value in (("Task", t["title"]), ("Notes", t["note"])):
            cell = ws[2][header.index(name)]
            self.assertEqual((cell.value, cell.data_type), (value, "s"), name)
        self.assertEqual((wb["Notes"]["A2"].value, wb["Notes"]["A2"].data_type), ("=2+3", "s"))
