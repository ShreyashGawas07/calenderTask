"""Export everything to an .xlsx workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import model
from .dialogs import fmt_date, repeat_label


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2A4A5C")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    return ws


def checklist_text(t, day):
    """'1/3: a ✓; b; c' (ticks for the task's own day, or today for a repeating task)."""
    items = t.get("items") or []
    if not items:
        return ""
    done, total = model.progress(t, day)
    return f"{done}/{total}: " + "; ".join(i["text"] + (" ✓" if model.item_done(t, i, day) else "") for i in items)


def export_xlsx(data, path, today):
    wb = Workbook()
    wb.remove(wb.active)
    tasks = sorted(data["tasks"], key=lambda t: (t["date"], t.get("time") or "99:99"))
    _sheet(wb, "Tasks",
           ["Date", "Time", "Task", "Checklist", "Notes", "Status", "Repeats", "Until", "Reminder (min before)"],
           [[fmt_date(model.s2d(t["date"])), t.get("time") or "", t["title"],
             checklist_text(t, model.s2d(t["date"]) if t.get("repeat", "none") == "none" else today),
             t.get("note", ""),
             ("Done" if t.get("done") else "Open") if t.get("repeat", "none") == "none"
             else f"{len(t.get('done_dates', []))} done",
             repeat_label(t), fmt_date(model.s2d(t.get("until"))),
             t.get("remind_before", 0) if t.get("remind") else ""] for t in tasks],
           [12, 8, 40, 40, 40, 12, 24, 12, 12])
    cds = []
    for c in data["countdowns"]:
        days, target, years = model.countdown_info(c, today)
        cds.append([c["title"], fmt_date(model.s2d(c["date"])), "Yes" if c.get("yearly") else "No",
                    fmt_date(target), days])
    _sheet(wb, "Countdowns", ["Title", "Date", "Every year", "Next", "Days left"], cds, [30, 12, 10, 12, 10])
    _sheet(wb, "Notes", ["Note"], [[n.get("text", "")] for n in data["notes"]], [80])
    wb.save(path)
